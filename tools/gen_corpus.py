#!/usr/bin/env python3
"""Generate the synthetic PO corpus with ground truth.

    python tools/gen_corpus.py --out tools/corpus

Produces 30 text PDFs (10 layouts × 3), 5 XLSX, 5 EML, 5 scanned-style PDFs,
one <doc_id>.json per document, index.json and catalog.json. Deliberate traps:
unknown codes, +6% price lines (and +1% inside tolerance), one duplicate PO
number, one totals mismatch, one sender on a webmail domain.
"""

from __future__ import annotations

import argparse
import datetime as dt
import io
import json
import logging
import random
from email.message import EmailMessage
from email.utils import format_datetime, make_msgid
from pathlib import Path

import demo_data

log = logging.getLogger("gen_corpus")
HERE = Path(__file__).resolve().parent
LAYOUT_DIR = HERE / "corpus" / "layouts"
LAYOUTS = [f"L{i:02d}" for i in range(1, 11)]

# Per layout: number style (None = partner's own), date format, show code column.
LAYOUT_STYLE = {
    "L01": (None, "%Y-%m-%d", True),
    "L02": ("us", "%m/%d/%Y", True),
    "L03": ("se", "%Y-%m-%d", True),
    "L04": ("eu", "%d.%m.%Y", True),
    "L05": (None, "%b %d, %Y", True),
    "L06": ("eu", "%d.%m.%Y", True),
    "L07": (None, "%Y-%m-%d", True),
    "L08": ("us", "%d-%b-%Y", True),
    "L09": (None, "%d %B %Y", False),
    "L10": ("us", "%Y%m%d", True),
}

# Trap plan. Keys are doc ids; values are sets of trap names.
PDF_TRAPS = {
    "pdf-L02-2": {"unknown"},
    "pdf-L05-1": {"unknown"},
    "pdf-L07-3": {"unknown", "deviation1"},
    "pdf-L03-3": {"deviation6"},
    "pdf-L08-1": {"deviation6", "deviation1"},
    "pdf-L06-3": {"mismatch"},
    "pdf-L04-2": {"duplicate"},  # same partner + PO number as pdf-L04-1
}
XLSX_TRAPS = {"xlsx-3": {"unknown"}, "xlsx-5": {"deviation6"}}
EML_TRAPS = {"eml-2": {"deviation6"}, "eml-4": {"webmail"}}
# Language-specific layouts only go to partners for whom they are plausible.
LAYOUT_PARTNERS = {"L03": [0, 6, 0], "L06": [2, 5, 2]}
SCAN_LAYOUTS = ["L01", "L03", "L05", "L08", "L10"]
XLSX_PARTNERS = [1, 4, 0, 6, 2]
EML_PARTNERS = [3, 5, 1, 7, 2]
SCAN_PARTNERS = [0, 2, 4, 7, 3]
QUANTITIES = [2, 4, 5, 6, 10, 12, 20, 24, 25, 30, 40, 48, 50, 60, 100, 120, 200, 250, 500]


# ----------------------------------------------------------------- formatting
def fmt_num(value: float, style: str, decimals: int = 2) -> str:
    s = f"{value:,.{decimals}f}"
    if style == "eu":
        return s.replace(",", "\x00").replace(".", ",").replace("\x00", ".")
    if style == "se":
        return s.replace(",", "\xa0").replace(".", ",")
    return s


def fmt_qty(value: float, style: str) -> str:
    if float(value).is_integer():
        return fmt_num(value, style, 0)
    return fmt_num(value, style, 1)


# ------------------------------------------------------------------- orders
class Generator:
    def __init__(self, out: Path, seed: int = demo_data.SEED):
        self.out = out
        self.rng = random.Random(seed + 1)
        self.catalog = demo_data.build_catalog(seed)
        self.products = self.catalog["products"]
        self.by_code = {p["default_code"]: p for p in self.products}
        self.sku_map = {}
        for row in self.catalog["sku_map"]:
            self.sku_map.setdefault(row["partner"], {})[row["product_code"]] = row["customer_code"]
        self.po_seq = {p["key"]: self.rng.randint(100, 900) for p in self.catalog["partners"]}
        self.docs: list[dict] = []

    # -- building blocks -----------------------------------------------------
    def partner(self, i: int) -> dict:
        return self.catalog["partners"][i % 8]

    def next_po(self, partner: dict) -> str:
        self.po_seq[partner["key"]] += self.rng.randint(1, 9)
        return partner["po_style"].format(y=2026, n=self.po_seq[partner["key"]])

    def pick_products(self, partner: dict, n: int) -> list[dict]:
        mapped = self.sku_map.get(partner["key"], {})
        pool_mapped = [self.by_code[c] for c in mapped]
        n_mapped = min(len(pool_mapped), round(n * 0.7)) if mapped else 0
        chosen = self.rng.sample(pool_mapped, n_mapped) if n_mapped else []
        rest = [p for p in self.products if p not in chosen]
        chosen += self.rng.sample(rest, n - n_mapped)
        self.rng.shuffle(chosen)
        return chosen

    def line(self, partner: dict, product: dict, show_codes: bool, trap: str | None) -> dict:
        pl_price = demo_data.pricelist_price(product, partner)
        unit_price = pl_price
        if trap == "deviation6":
            unit_price = round(pl_price * 1.06, 2)
        elif trap == "deviation1":
            unit_price = round(pl_price * 1.01, 2)
        deviation = round((unit_price - pl_price) / pl_price * 100, 2) if pl_price else 0.0
        if product["uom"] in ("m", "kg") and self.rng.random() < 0.3:
            qty = self.rng.choice([2.5, 7.5, 12.5, 37.5])
        else:
            qty = self.rng.choice(QUANTITIES)
        mapped = self.sku_map.get(partner["key"], {}).get(product["default_code"])
        if not show_codes:
            code, method = None, "fuzzy"
        elif mapped:
            code, method = mapped, "sku_map"
        else:
            code, method = product["default_code"], "default_code"
        return {
            "customer_code": code,
            "description": product["name"],
            "quantity": qty,
            "uom": product["uom"],
            "unit_price": unit_price,
            "line_total": round(qty * unit_price, 2),
            "_product_code": product["default_code"],
            "_match": method,
            "_pricelist_price": pl_price,
            "_deviation_pct": deviation,
            "_price_flag": abs(deviation) > 2.0,
        }

    def unknown_line(self, partner: dict) -> dict:
        prefix = partner["own_codes"] or "ART"
        code = f"{prefix}-{self.rng.randint(90000, 99999)}"
        family = self.rng.choice(demo_data.PRODUCT_FAMILIES)
        desc = f"{self.rng.choice(family[1])} {self.rng.choice(family[2])} special"
        qty = self.rng.choice(QUANTITIES)
        price = round(self.rng.uniform(1.5, 60.0), 2)
        return {
            "customer_code": code,
            "description": desc,
            "quantity": qty,
            "uom": family[4],
            "unit_price": price,
            "line_total": round(qty * price, 2),
            "_product_code": None,
            "_match": "none",
            "_pricelist_price": None,
            "_deviation_pct": None,
            "_price_flag": False,
        }

    def order(
        self,
        partner: dict,
        n_lines: int,
        traps: set[str],
        show_codes: bool = True,
        po_number: str | None = None,
    ) -> dict:
        products = self.pick_products(partner, n_lines)
        # Deviation lines need a price high enough that rounding keeps the % intact.
        rich = [i for i, p in enumerate(products) if demo_data.pricelist_price(p, partner) >= 20]
        slots = list(rich) or list(range(n_lines))
        self.rng.shuffle(slots)
        line_traps: dict[int, str] = {}
        for trap in ("deviation6", "deviation1"):
            if trap in traps and slots:
                line_traps[slots.pop()] = trap
        lines = [
            self.line(partner, p, show_codes, line_traps.get(i)) for i, p in enumerate(products)
        ]
        if "unknown" in traps:
            lines.insert(self.rng.randrange(len(lines) + 1), self.unknown_line(partner))
        po_date = dt.date(2026, 9, 1) + dt.timedelta(days=self.rng.randint(0, 50))
        requested = po_date + dt.timedelta(days=self.rng.randint(7, 21))
        subtotal = round(sum(ln["line_total"] for ln in lines), 2)
        tax = round(subtotal * partner["vat_rate"], 2)
        total = round(subtotal + tax, 2)
        if "mismatch" in traps:
            # Clearly outside the totals tolerance (0.5 % or one unit): +5 % and +100.
            total = round(total * 1.05 + 100.0, 2)
        buyer = partner["contacts"][0]
        notes = self.rng.choice(
            [
                None,
                "Please confirm delivery date.",
                "Deliver to goods receiving, gate B.",
                "Partial delivery accepted.",
            ]
        )
        return {
            "po_number": po_number or self.next_po(partner),
            "po_date": po_date.isoformat(),
            "requested_date": requested.isoformat(),
            "currency": partner["currency"],
            "buyer_name": partner["name"],
            "buyer_email": buyer["email"],
            "buyer_vat": partner["vat"] or None,
            "ship_to": (
                f"{partner['name']}, {partner['street']}, {partner['zip']} {partner['city']}"
            ),
            "lines": lines,
            "subtotal": subtotal,
            "tax": tax,
            "total": total,
            "notes": notes,
        }

    # -- ground truth ---------------------------------------------------------
    def record(
        self,
        doc_id: str,
        source_type: str,
        layout: str | None,
        file: str,
        partner: dict | None,
        sender: str,
        parsed: dict,
        duplicate_of: str | None,
        mismatch: bool,
    ) -> dict:
        expected_lines = []
        for ln in parsed["lines"]:
            expected_lines.append(
                {
                    "customer_code": ln["customer_code"],
                    "product_code": ln["_product_code"],
                    "match_method": ln["_match"],
                    "pricelist_price": ln["_pricelist_price"],
                    "deviation_pct": ln["_deviation_pct"],
                    "price_flag": ln["_price_flag"],
                }
            )
        clean_lines = [
            {k: v for k, v in ln.items() if not k.startswith("_")} for ln in parsed["lines"]
        ]
        unknown = sum(1 for ln in expected_lines if ln["match_method"] == "none")
        price_flags = sum(1 for ln in expected_lines if ln["price_flag"])
        doc = {
            "doc_id": doc_id,
            "source_type": source_type,
            "layout": layout,
            "file": file,
            "sender_email": sender,
            "parsed": {**parsed, "lines": clean_lines},
            "expected": {
                "partner": partner["key"] if partner else None,
                "partner_resolution": "email" if partner else "none",
                "lines": expected_lines,
                "unknown_codes": unknown,
                "price_flags": price_flags,
                "duplicate_of": duplicate_of,
                "totals_mismatch": mismatch,
                "stop_expected": bool(
                    unknown or mismatch or price_flags or duplicate_of or partner is None
                ),
            },
        }
        (self.out / f"{doc_id}.json").write_text(json.dumps(doc, indent=2, ensure_ascii=False))
        self.docs.append(doc)
        return doc

    # -- renderers --------------------------------------------------------------
    @staticmethod
    def logo_svg(partner: dict, height: int = 34) -> str:
        """A deterministic inline-SVG logo per partner: coloured mark + wordmark.
        Real POs carry the buyer's logo; extraction must ignore it."""
        import hashlib

        h = hashlib.sha1(partner["key"].encode()).digest()
        hue = h[0] * 360 // 255
        color = f"hsl({hue}, 55%, 38%)"
        initials = "".join(w[0] for w in partner["name"].split()[:2]).upper()
        style = h[1] % 3
        name = partner["name"]
        if style == 0:  # square monogram + name
            mark = (
                f'<rect width="{height}" height="{height}" rx="4" fill="{color}"/>'
                f'<text x="{height / 2}" y="{height * 0.68}" font-size="{height * 0.5}" '
                f'font-family="Helvetica, Arial" font-weight="bold" fill="#fff" '
                f'text-anchor="middle">{initials}</text>'
            )
        elif style == 1:  # circle + name
            mark = (
                f'<circle cx="{height / 2}" cy="{height / 2}" r="{height / 2}" fill="{color}"/>'
                f'<text x="{height / 2}" y="{height * 0.68}" font-size="{height * 0.45}" '
                f'font-family="Helvetica, Arial" font-weight="bold" fill="#fff" '
                f'text-anchor="middle">{initials}</text>'
            )
        else:  # three bars + name
            w = height * 0.22
            mark = "".join(
                f'<rect x="{i * (w + 3)}" y="{height * (0.5 - 0.5 * s)}" width="{w}" '
                f'height="{height * s}" rx="2" fill="{color}"/>'
                for i, s in enumerate((0.6, 1.0, 0.8))
            )
        width = height + 12 + int(len(name) * height * 0.58)
        return (
            f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" '
            f'viewBox="0 0 {width} {height}">{mark}'
            f'<text x="{height + 10}" y="{height * 0.66}" font-size="{height * 0.5}" '
            f'font-family="Helvetica, Arial" font-weight="bold" fill="{color}">{name}</text></svg>'
        )

    def render_html(self, layout: str, partner: dict, parsed: dict) -> str:
        from jinja2 import Environment, FileSystemLoader
        from markupsafe import Markup

        num_style, date_fmt, show_codes = LAYOUT_STYLE[layout]
        style = num_style or partner["numbers"]
        env = Environment(loader=FileSystemLoader(str(LAYOUT_DIR)), autoescape=True)

        def date(s):
            return dt.date.fromisoformat(s).strftime(date_fmt)

        lines = []
        for i, ln in enumerate(parsed["lines"], 1):
            lines.append(
                {
                    "pos": i,
                    "code": ln["customer_code"] or "",
                    "desc": ln["description"],
                    "qty": fmt_qty(ln["quantity"], style),
                    "uom": ln["uom"],
                    "price": fmt_num(ln["unit_price"], style),
                    "total": fmt_num(ln["line_total"], style),
                    "delivery": date(parsed["requested_date"]),
                }
            )
        ctx = {
            "company": self.catalog["company"],
            "partner": partner,
            "po": parsed,
            "po_date": date(parsed["po_date"]),
            "requested_date": date(parsed["requested_date"]),
            "lines": lines,
            "subtotal": fmt_num(parsed["subtotal"], style),
            "tax": fmt_num(parsed["tax"], style),
            "tax_pct": f"{partner['vat_rate'] * 100:.0f}",
            "total": fmt_num(parsed["total"], style),
            "cur": parsed["currency"],
            "show_codes": show_codes,
            "logo": Markup(self.logo_svg(partner)),
            "logo_big": Markup(self.logo_svg(partner, 44)),
        }
        return env.get_template(f"{layout}.html").render(**ctx)

    def render_pdf(self, html: str, path: Path) -> None:
        from weasyprint import HTML

        HTML(string=html, base_url=str(LAYOUT_DIR)).write_pdf(str(path))

    def render_scan(self, html: str, path: Path) -> None:
        """Rasterise the rendered PDF, add noise/skew/JPEG artefacts, save image-only PDF."""
        import numpy as np
        import pypdfium2 as pdfium
        from PIL import Image, ImageFilter
        from weasyprint import HTML

        pdf_bytes = HTML(string=html, base_url=str(LAYOUT_DIR)).write_pdf()
        pdf = pdfium.PdfDocument(pdf_bytes)
        pages = []
        for page in pdf:
            img = page.render(scale=130 / 72).to_pil().convert("L")
            img = img.rotate(self.rng.uniform(-0.8, 0.8), resample=Image.BICUBIC, fillcolor=235)
            arr = np.asarray(img).astype(np.float32)
            arr = arr * 0.93 + 10  # grey-ish paper
            arr += np.random.default_rng(self.rng.randint(0, 2**31)).normal(0, 9, arr.shape)
            img = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))
            img = img.filter(ImageFilter.GaussianBlur(0.7))
            buf = io.BytesIO()
            img.save(buf, "JPEG", quality=55)
            pages.append(Image.open(io.BytesIO(buf.getvalue())).convert("L"))
        pages[0].save(str(path), "PDF", save_all=True, append_images=pages[1:], resolution=130)

    def render_xlsx(self, variant: int, partner: dict, parsed: dict, path: Path) -> None:
        import openpyxl
        from openpyxl.styles import Font

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Lines" if variant == 4 else "PO"
        headers = {
            1: ["Code", "Description", "Qty", "Unit", "Unit price", "Total"],
            2: ["Item no.", "Description", "Quantity", "UoM", "Price", "Amount"],
            3: ["Description", "Our code", "Qty", "Unit price", "Line total", "Notes"],
            4: ["SKU", "Product", "Qty", "Unit", "Price", "Total"],
            5: ["Art.nr", "Benämning", "Antal", "Enhet", "À-pris", "Summa"],
        }[variant]
        meta = [
            ("PO number", parsed["po_number"]),
            ("Date", parsed["po_date"]),
            ("Requested delivery", parsed["requested_date"]),
            ("Buyer", partner["name"]),
            ("Currency", parsed["currency"]),
        ]
        if variant == 4:
            info = wb.create_sheet("Info", 0)
            for r, (k, v) in enumerate(meta, 1):
                info.cell(r, 1, k)
                info.cell(r, 2, v)
            start = 1
        elif variant == 1:
            for c, (k, v) in enumerate(meta, 1):
                ws.cell(1, c, f"{k}: {v}")
            start = 3
        else:
            ws.merge_cells("A1:F1")
            ws.cell(1, 1, f"PURCHASE ORDER {parsed['po_number']}").font = Font(bold=True, size=14)
            for r, (k, v) in enumerate(meta, 2):
                ws.cell(r, 1, k)
                ws.cell(r, 2, v)
            start = len(meta) + 3
        for c, h in enumerate(headers, 1):
            ws.cell(start, c, h).font = Font(bold=True)
        for r, ln in enumerate(parsed["lines"], start + 1):
            if variant == 3:
                row = [
                    ln["description"],
                    ln["customer_code"],
                    ln["quantity"],
                    ln["unit_price"],
                    ln["line_total"],
                    "",
                ]
            else:
                row = [
                    ln["customer_code"],
                    ln["description"],
                    ln["quantity"],
                    ln["uom"],
                    ln["unit_price"],
                    ln["line_total"],
                ]
            for c, v in enumerate(row, 1):
                ws.cell(r, c, v)
        r = start + len(parsed["lines"]) + 2
        for label, val in (
            ("Subtotal", parsed["subtotal"]),
            ("VAT", parsed["tax"]),
            ("Total", parsed["total"]),
        ):
            ws.cell(r, 5, label).font = Font(bold=True)
            ws.cell(r, 6, val)
            r += 1
        wb.save(str(path))

    def render_eml(
        self, variant: int, partner: dict, parsed: dict, sender: dict, path: Path
    ) -> None:
        style = partner["numbers"]
        body = ["Hello,", "", f"Please find our purchase order {parsed['po_number']} below.", ""]
        for ln in parsed["lines"]:
            qty = fmt_qty(ln["quantity"], style)
            price = fmt_num(ln["unit_price"], style)
            if variant in (1, 4):
                body.append(
                    f"- {qty} {ln['uom']} {ln['customer_code']} {ln['description']}"
                    f" @ {price} {parsed['currency']}"
                )
            elif variant == 2:
                body.append(
                    f"{ln['customer_code']:<12} {ln['description']:<44} {qty:>7} x {price:>10}"
                )
            else:
                body.append(
                    f"{qty} x {ln['description']} (ref {ln['customer_code']}) - unit price {price}"
                )
        body += [
            "",
            f"Subtotal: {fmt_num(parsed['subtotal'], style)} {parsed['currency']}",
            f"Total incl. VAT: {fmt_num(parsed['total'], style)} {parsed['currency']}",
            f"Requested delivery: {parsed['requested_date']}",
            "",
            "Best regards,",
            sender["name"],
            partner["name"],
            f"{partner['street']}, {partner['zip']} {partner['city']}",
            "",
            "-----Original Message-----",
            f"From: {self.catalog['company']['email']}",
            "Subject: Re: price list 2026",
            "> Thanks for the updated price list, we will send the next order shortly.",
            "> Regards",
        ]
        msg = EmailMessage()
        msg["From"] = f"{sender['name']} <{sender['email']}>"
        msg["To"] = self.catalog["company"]["email"]
        msg["Subject"] = f"Purchase order {parsed['po_number']} - {partner['name']}"
        msg["Date"] = format_datetime(
            dt.datetime.fromisoformat(parsed["po_date"] + "T09:30:00+02:00")
        )
        msg["Message-ID"] = make_msgid(domain=sender["email"].split("@")[1])
        text = "\n".join(body)
        msg.set_content(text)
        if variant == 5:
            html = "<html><body>" + "".join(f"<p>{line}</p>" for line in body) + "</body></html>"
            msg.add_alternative(html, subtype="html")
        path.write_bytes(msg.as_bytes())

    # -- the plan ---------------------------------------------------------------
    def run(self) -> list[dict]:
        self.out.mkdir(parents=True, exist_ok=True)
        (self.out / "catalog.json").write_text(
            json.dumps(self.catalog, indent=2, ensure_ascii=False)
        )
        self.docs = []
        dup_source: dict | None = None
        pi = 0
        for layout in LAYOUTS:
            for inst in (1, 2, 3):
                doc_id = f"pdf-{layout}-{inst}"
                if layout in LAYOUT_PARTNERS:
                    partner = self.partner(LAYOUT_PARTNERS[layout][inst - 1])
                else:
                    partner = self.partner(pi)
                    pi += 1
                traps = set(PDF_TRAPS.get(doc_id, set()))
                n_lines = 38 if layout == "L07" else self.rng.randint(5, 14)
                show_codes = LAYOUT_STYLE[layout][2]
                po_number, duplicate_of = None, None
                if "duplicate" in traps and dup_source:
                    partner = dup_source["partner"]
                    po_number = dup_source["parsed"]["po_number"]
                    duplicate_of = dup_source["doc_id"]
                parsed = self.order(partner, n_lines, traps, show_codes, po_number)
                html = self.render_html(layout, partner, parsed)
                self.render_pdf(html, self.out / f"{doc_id}.pdf")
                self.record(
                    doc_id,
                    "pdf",
                    layout,
                    f"{doc_id}.pdf",
                    partner,
                    partner["contacts"][0]["email"],
                    parsed,
                    duplicate_of,
                    "mismatch" in traps,
                )
                if doc_id == "pdf-L04-1":
                    dup_source = {"partner": partner, "parsed": parsed, "doc_id": doc_id}
                log.info("wrote %s (%s, %d lines)", doc_id, partner["key"], len(parsed["lines"]))
        for i, p_idx in enumerate(XLSX_PARTNERS, 1):
            doc_id = f"xlsx-{i}"
            partner = self.partner(p_idx)
            traps = XLSX_TRAPS.get(doc_id, set())
            parsed = self.order(partner, self.rng.randint(6, 16), traps)
            self.render_xlsx(i, partner, parsed, self.out / f"{doc_id}.xlsx")
            self.record(
                doc_id,
                "xlsx",
                f"X{i}",
                f"{doc_id}.xlsx",
                partner,
                partner["contacts"][1]["email"],
                parsed,
                None,
                False,
            )
        for i, p_idx in enumerate(EML_PARTNERS, 1):
            doc_id = f"eml-{i}"
            partner = self.partner(p_idx)
            traps = EML_TRAPS.get(doc_id, set())
            parsed = self.order(partner, self.rng.randint(3, 8), traps)
            sender = partner["contacts"][-1] if "webmail" in traps else partner["contacts"][0]
            self.render_eml(i, partner, parsed, sender, self.out / f"{doc_id}.eml")
            resolved = None if "webmail" in traps else partner
            self.record(
                doc_id,
                "eml",
                f"E{i}",
                f"{doc_id}.eml",
                resolved,
                sender["email"],
                parsed,
                None,
                False,
            )
        for i, (layout, p_idx) in enumerate(zip(SCAN_LAYOUTS, SCAN_PARTNERS, strict=False), 1):
            doc_id = f"scan-{i}"
            partner = self.partner(p_idx)
            parsed = self.order(partner, self.rng.randint(5, 10), set())
            self.render_scan(self.render_html(layout, partner, parsed), self.out / f"{doc_id}.pdf")
            self.record(
                doc_id,
                "scan",
                layout,
                f"{doc_id}.pdf",
                partner,
                partner["contacts"][0]["email"],
                parsed,
                None,
                False,
            )
        index = {
            "seed": self.catalog["seed"],
            "count": len(self.docs),
            "documents": [
                {k: d[k] for k in ("doc_id", "source_type", "layout", "file", "sender_email")}
                | {
                    "expected_partner": d["expected"]["partner"],
                    "stop_expected": d["expected"]["stop_expected"],
                }
                for d in self.docs
            ],
        }
        (self.out / "index.json").write_text(json.dumps(index, indent=2))
        return self.docs


def generate(out: Path, seed: int = demo_data.SEED) -> list[dict]:
    return Generator(Path(out), seed).run()


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--out", default=str(HERE / "corpus"), help="output directory")
    ap.add_argument("--seed", type=int, default=demo_data.SEED)
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args()
    logging.basicConfig(
        level=logging.INFO if args.verbose else logging.WARNING, format="%(message)s"
    )
    docs = generate(Path(args.out), args.seed)
    counts = {}
    for d in docs:
        counts[d["source_type"]] = counts.get(d["source_type"], 0) + 1
    stops = sum(1 for d in docs if d["expected"]["stop_expected"])
    print(f"{len(docs)} documents in {args.out}: {counts}; {stops} expected to stop")


if __name__ == "__main__":
    main()
